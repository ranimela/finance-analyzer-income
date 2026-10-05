"""Tests for reconciliation working file generator."""

import csv
from pathlib import Path

from src.reconciliation import generate_income_working_file


def test_generate_income_working_file(tmp_path: Path) -> None:
    """Test generating income_working.csv with occurrences and date resolution."""
    ledger_csv = Path("data/outputs/כרטסות הכנסות 24-26_step9.csv")
    ap_xlsx = Path("data/inputs/Agreements Payments.xlsx")

    assert ledger_csv.exists()
    assert ap_xlsx.exists()

    output_csv = tmp_path / "test_income_working.csv"
    metrics = generate_income_working_file(ledger_csv, ap_xlsx, output_csv)

    # 1. Check metric consistency
    assert metrics["total_ledger_rows"] == 1367
    assert metrics["zero_or_non_positive_rows"] == 30
    assert metrics["ap_zero_matches"] == 506
    assert metrics["ap_single_matches"] == 392
    assert metrics["ap_multiple_matches_consistent"] == 391
    assert metrics["ap_multiple_matches_has_error"] == 48
    assert metrics["start_date_errors"] == 45
    assert metrics["end_date_errors"] == 46
    assert metrics["duration_calculated_rows"] == 666
    assert metrics["duration_skipped_rows"] == 464
    assert metrics["adjusted_income_calculated_rows"] == 593
    assert metrics["special_mion_direct_allocation_rows"] == 237
    assert metrics["quarter_columns_count"] == 31
    assert metrics["max_allocated_quarter"] == "31Q3"
    assert metrics["rows_with_quarter_allocations"] == 773
    assert metrics["total_output_rows"] == 1367

    # 2. Check output file integrity
    assert output_csv.exists()
    with open(output_csv, mode="r", encoding="utf-8-sig", newline="") as f:
        reader = list(csv.reader(f))

    assert len(reader) == 1368  # 1 header + 1367 data rows
    header = reader[0]
    assert len(header) == 46  # 15 base/calculated + 31 quarter columns (24Q1 to 31Q3)

    assert header[15] == "24Q1"
    assert header[-1] == "31Q3"

    # Verify duration and adjusted monthly income calculation logic on first data row
    first_row = reader[1]
    assert first_row[13] == "3"  # 01/07/2026 to 30/09/2026 -> 3 months
    assert first_row[14] == "16271.00"  # 48813 / 3 = 16271.00
    col_26q3_idx = header.index("26Q3")
    assert first_row[col_26q3_idx] == "48813.00"

    # Verify fallback lookup on Col F=0 rows (e.g. key 250706 in Col G)
    fallback_sample = [r for r in reader[1:] if r[5] in ("0", "0.0") and r[6] == "250706" and r[9] != ""][0]
    assert fallback_sample[10] == "2"  # 2 occurrences found in AP
    assert fallback_sample[11] == "01/08/2025"
    assert fallback_sample[12] == "31/10/2025"
    assert fallback_sample[13] == "3"

    # Verify direct quarter assignment on 93002 row
    special_sample = [r for r in reader[1:] if r[2] == "93002" and r[3] == "2024-06-30"][0]
    assert special_sample[13] == ""  # duration skipped
    assert special_sample[14] == ""  # adjusted income skipped
    col_24q2_idx = header.index("24Q2")
    assert special_sample[col_24q2_idx] == "6233164.75"

    # 3. Check error flagging behavior
    # Invoice Start Date is index 11, Invoice End Date is index 12
    start_err_rows = [r for r in reader[1:] if r[11] == "ERROR - multiple dates"]
    end_err_rows = [r for r in reader[1:] if r[12] == "ERROR - multiple dates"]
    assert len(start_err_rows) == 45
    assert len(end_err_rows) == 46

    # Verify dd/mm/yyyy format on valid single matches
    single_match_rows = [r for r in reader[1:] if r[10] == "1"]
    assert len(single_match_rows) == 392
    for r in single_match_rows:
        s_val, e_val = r[11], r[12]
        if s_val:
            assert len(s_val) == 10 and s_val[2] == "/" and s_val[5] == "/", f"Invalid format: {s_val}"
        if e_val:
            assert len(e_val) == 10 and e_val[2] == "/" and e_val[5] == "/", f"Invalid format: {e_val}"


def test_process_cancellations(tmp_path: Path) -> None:
    """Test process_cancellations cuts both cancellation notices and matched original invoices."""
    from src.reconciliation import process_cancellations

    input_csv = Path("data/outputs/income_working.csv")
    assert input_csv.exists()

    output_csv = tmp_path / "income_working_step2.csv"
    output_canc_csv = tmp_path / "cancellations.csv"
    output_xlsx = tmp_path / "income_working_step2.xlsx"

    metrics = process_cancellations(
        input_csv,
        output_csv,
        output_cancellations_csv_path=output_canc_csv,
        output_step2_xlsx_path=output_xlsx,
    )

    assert metrics["total_input_rows"] == 1367
    assert metrics["cancellation_notice_rows"] == 54
    assert metrics["target_invoices_extracted"] == 49
    assert metrics["original_invoice_rows_matched_in_col_f"] == 56
    assert metrics["total_cancellations_cut"] == 109
    assert metrics["retained_step2_rows"] == 1258

    # Verify CSV line counts
    with open(output_csv, mode="r", encoding="utf-8-sig") as f:
        step2_rows = list(csv.reader(f))
    assert len(step2_rows) == 1259  # 1 header + 1258 data rows

    with open(output_canc_csv, mode="r", encoding="utf-8-sig") as f:
        canc_rows = list(csv.reader(f))
    assert len(canc_rows) == 110  # 1 header + 109 data rows

    # Verify Excel workbook tabs
    import openpyxl
    wb = openpyxl.load_workbook(output_xlsx)
    assert wb.sheetnames == ["income_working_step2", "Cancellations"]
    assert wb["income_working_step2"].max_row == 1259
    assert wb["Cancellations"].max_row == 110
